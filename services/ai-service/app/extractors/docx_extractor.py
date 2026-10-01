"""DOCX extractor using python-docx when installed."""

from __future__ import annotations

import io

from app.extractors.base import BaseExtractor, ExtractionResult, unavailable_result
from app.logging_config import get_logger

logger = get_logger(__name__)


class DocxExtractor(BaseExtractor):
    """Extract paragraphs and tables from a .docx document."""

    name = "docx"

    def supports(self, mime_type: str | None, extension: str | None) -> bool:
        ext = (extension or "").lstrip(".").lower()
        mime = (mime_type or "").lower()
        return ext in {"docx", "docm"} or mime in {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.template",
        }

    def extract(
        self,
        data: bytes,
        *,
        filename: str | None = None,
        mime_type: str | None = None,
        extension: str | None = None,
    ) -> ExtractionResult:
        try:
            import docx
        except ImportError:
            logger.warning("python-docx not installed; cannot extract DOCX")
            return unavailable_result(
                mime_type,
                "python-docx is not installed",
                extraction_method="docx-unavailable",
            )

        warnings: list[str] = []
        try:
            document = docx.Document(io.BytesIO(data))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text for cell in row.cells if cell.text]
                    if cells:
                        parts.append("\t".join(cells))
            text = "\n".join(parts).strip()
            return ExtractionResult(
                text=text,
                extraction_method=self.name,
                mime_type=mime_type,
                page_count=None,
                warnings=warnings,
            )
        except Exception as exc:  # noqa: BLE001 - corrupt docs must not 500
            logger.warning("python-docx failed: %s", exc)
            return unavailable_result(
                mime_type,
                f"failed to read DOCX: {exc}",
                extraction_method="docx-error",
            )
